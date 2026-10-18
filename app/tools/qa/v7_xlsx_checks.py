"""Read generated fixture exports; no connection to a working database."""
from pathlib import Path
import io,json,posixpath,zipfile,datetime
import xml.etree.ElementTree as ET
from PIL import Image
from openpyxl import load_workbook

out=Path(__file__).parent/'output/v7'
checks=[]
for path in sorted(out.glob('*.xlsx')):
    archive=zipfile.ZipFile(path)
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith(('.xml','.rels','.svg')):
            ET.fromstring(archive.read(name))
        if name.endswith('.rels'):
            base=posixpath.dirname(name).replace('/_rels','') if name!='_rels/.rels' else ''
            for rel in ET.fromstring(archive.read(name)):
                if rel.get('TargetMode')!='External':
                    target=posixpath.normpath(posixpath.join(base,rel.get('Target')))
                    assert target in archive.namelist(),(name,target)
    book=load_workbook(path,data_only=False)
    assert book.sheetnames
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type!='e',(path,cell.coordinate,cell.value)
    if path.name=='technology.xlsx':
        sheet=book.worksheets[0]
        expected=['Этап №','Этап','Операция ID','Операция','Тип','Оборудование','Партия','Секунды','Оценка, сек','Подготовка, мин','Разряд','Вариант каталога','Проверено','Зависимости','Примечания']
        assert [c.value for c in sheet[4]]==expected
        assert sheet.freeze_panes=='A5' and str(sheet.merged_cells)=='A1:O1 A2:O2'
        assert sheet['H6'].data_type=='n' and sheet['I6'].value is None
        assert 'xl/media/garment.svg' in archive.namelist()
        png=Image.open(io.BytesIO(archive.read('xl/media/garment.png')))
        png.verify()
        assert len(book['Схема изделия']._images)==1
    if path.name=='goals.xlsx':
        assert isinstance(book.worksheets[0]['F5'].value,datetime.datetime)
        assert book.worksheets[0]['L5'].data_type=='n'
    if path.name=='formula-literal.xlsx':
        assert book.active['A5'].data_type=='s' and book.active['A5'].value.startswith('=')
    if path.name=='worker-production.xlsx':
        for row in book.worksheets[0].iter_rows(min_row=5,values_only=True):
            assert 'p2' not in row
    checks.append({'file':path.name,'sheets':len(book.sheetnames),'bytes':path.stat().st_size})
report={'passed':len(checks),'files':checks}
(out/'xlsx-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False))
