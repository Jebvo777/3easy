"""Extract technical fields only; never publish employee assignments or piece rates.
Run with the bundled Python: script CATALOG.xlsx EXAMPLES_DIRECTORY.
This prepares data, not application QA.
"""
import sys, json, re, hashlib
from pathlib import Path
import openpyxl, xlrd

root = Path(__file__).resolve().parents[1]
out = root / 'app/data'
out.mkdir(exist_ok=True)
catalog = Path(sys.argv[1])
w = openpyxl.load_workbook(catalog, read_only=True, data_only=True)
records = {}
for row in w['Номенклатура'].iter_rows(min_row=8, values_only=True):
    if not row[0]: continue
    records[row[0]] = dict(id=row[0], node=row[1], name=row[2], action=row[3], note=row[6], variants=[])
for row in w['Варианты'].iter_rows(min_row=8, values_only=True):
    if row[0] not in records: continue
    records[row[0]]['variants'].append(dict(id=row[1], construction=row[3], condition=row[4], seconds=row[5], method=row[6], grade=row[7], kind=row[8], source_ids=str(row[9] or '')))
(out/'operations.json').write_text(json.dumps({'schema':1,'source':'Выгрузка ТехОперации 1С 7.7 от 27.09.2026; условия применимости проверяет технолог','operations':list(records.values())},ensure_ascii=False,separators=(',',':')))

examples=[]
for p in sorted(Path(sys.argv[2]).iterdir()):
    if p.suffix.lower() not in ('.xlsx','.xls'): continue
    if p.suffix=='.xlsx':
        book=openpyxl.load_workbook(p,read_only=True,data_only=True)
        sheets=[(s.title,list(s.iter_rows(values_only=True))) for s in book]
    else:
        book=xlrd.open_workbook(str(p),encoding_override='cp1251',logfile=open('/dev/null','w'))
        sheets=[(s.name,[s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]
    preferred=next((s for s in sheets if s[0]=='последовательность'),sheets[0])
    ops=[];seen=set();node='Подготовка';node_n=1
    for row in preferred[1]:
        found=False
        for k in range(min(3,len(row)-1)):
            n,name=row[k],row[k+1]
            if isinstance(n,(int,float)) and int(n)==n and 1<=n<=1000 and isinstance(name,str) and len(name)>9 and re.search('[а-яА-Я]',name):
                if int(n) in seen: found=True; break
                seen.add(int(n)); tail=row[k+2:]; method=next((str(v) for v in tail[:3] if isinstance(v,str) and v.strip()),'')
                seconds=next((int(v) for v in tail[1:] if isinstance(v,(float,int)) and 0<v<36000),None)
                typ='iron' if method.lower().startswith('у') else 'overlock' if 'овер' in method.lower() else 'sew' if method.lower().startswith('м') else 'hand'
                ops.append(dict(id='example-'+str(len(ops)+1),name=name.strip(),type=typ,node_id='node-'+str(node_n),node_name=node,seconds=None,estimated_seconds=seconds,setup_minutes=0,dependencies=['example-'+str(len(ops))] if ops else [],equipment=method,grade=None,source=p.name+' / '+preferred[0],batch=20,norm_reviewed=False,notes='Образец: проверить применимость и время для своего изделия'))
                found=True;break
        if not found:
            strings=[v.strip() for v in row if isinstance(v,str) and v.strip()]
            candidate=next((s for s in strings if re.match(r'^(Обработка|Сборка|Монтаж|Отделка|Подготовка|Раскрой)',s,re.I)),None)
            if candidate and len(candidate)<180: node=candidate;node_n+=1
    if ops: examples.append(dict(id='sample-'+str(len(examples)+1),name=p.stem,source_file=p.name,sheet=preferred[0],sha256=hashlib.sha256(p.read_bytes()).hexdigest(),route=ops))
(out/'examples.json').write_text(json.dumps({'schema':1,'notice':'Техническая редакция: удалены имена, распределения, зарплаты и расценки. Время — образец для проверки, не утверждённая норма.','examples':examples},ensure_ascii=False,separators=(',',':')))
